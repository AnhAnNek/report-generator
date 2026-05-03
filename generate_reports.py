from openpyxl import load_workbook
from datetime import datetime, timedelta
from pathlib import Path
from docx import Document

INPUT_FILE = "input/TranVanAn.xlsx"
OUTPUT_DIR = Path("Generate")
OUTPUT_DIR.mkdir(exist_ok=True)


# ========================
# Utils
# ========================
def excel_date_to_datetime(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, (int, float)):
        base_date = datetime(1899, 12, 30)
        return base_date + timedelta(days=int(value))

    if isinstance(value, str):
        value = value.strip()

        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                pass

    return None


def date_to_display(date_obj):
    return date_obj.strftime("%d/%m/%Y")


def date_to_folder(date_obj):
    return date_obj.strftime("%Y-%m-%d")


def get_week_start(date_obj):
    # Monday is weekday 0
    return date_obj - timedelta(days=date_obj.weekday())


def get_daily_dir(date_obj):
    week_start = get_week_start(date_obj)
    week_folder = f"week-{date_to_folder(week_start)}"
    day_folder = date_to_folder(date_obj)

    folder = OUTPUT_DIR / week_folder / day_folder
    folder.mkdir(parents=True, exist_ok=True)

    return folder


def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


# ========================
# Read Excel
# ========================
def read_wbs(wb):
    ws = wb["WBS"]
    tasks = {}

    for row in ws.iter_rows(min_row=6, values_only=True):
        task_id = clean_text(row[1])
        task_name = clean_text(row[2])

        if not task_id:
            continue

        tasks[task_id] = {
            "task_id": task_id,
            "task_name": task_name,
            "estimate_h": row[3],
            "actual_h": row[6],
            "percent": row[9],
            "note": clean_text(row[10]),
        }

    return tasks


def read_daily(wb):
    ws = wb["Report Daily"]

    date_cols = {}

    for col in range(4, ws.max_column + 1):
        date_val = ws.cell(row=1, column=col).value
        date_obj = excel_date_to_datetime(date_val)

        if date_obj:
            date_cols[col] = date_obj

    data = []

    for row in range(5, ws.max_row + 1):
        task_id = clean_text(ws.cell(row=row, column=2).value)
        task_name = clean_text(ws.cell(row=row, column=3).value)

        if not task_id:
            continue

        for col, date_obj in date_cols.items():
            content = clean_text(ws.cell(row=row, column=col).value)

            # Important:
            # Nếu ngày đó không có nội dung trong Excel thì bỏ qua,
            # không tạo folder/file cho ngày đó.
            if not content:
                continue

            data.append({
                "date_obj": date_obj,
                "date": date_to_display(date_obj),
                "folder_date": date_to_folder(date_obj),
                "task_id": task_id,
                "task_name": task_name,
                "content": content,
            })

    return data


# ========================
# Generate TXT Reports
# ========================
def generate_reports_by_date(daily_data, tasks):
    grouped_by_date = {}

    for item in daily_data:
        key = item["folder_date"]
        grouped_by_date.setdefault(key, []).append(item)

    sorted_dates = sorted(grouped_by_date.keys())

    for i, date_key in enumerate(sorted_dates):
        items = grouped_by_date[date_key]

        # Group by task_id to prevent duplicates
        task_dict = {}
        for item in items:
            tid = item["task_id"]
            task_dict[tid] = item  # Keep the last if duplicates
        items = list(task_dict.values())

        date_obj = items[0]["date_obj"]
        date_display = items[0]["date"]

        # Get next date for NEXT DAY
        next_date_key = sorted_dates[i + 1] if i + 1 < len(sorted_dates) else None
        if next_date_key:
            next_items = grouped_by_date[next_date_key]
            next_task_dict = {item["task_id"]: item for item in next_items}
            next_tasks = "\n".join([f"- {item['task_name']}" for item in next_task_dict.values()])
        else:
            next_tasks = "- N/A"

        daily_dir = get_daily_dir(date_obj)

        # Prepare difficulties for morning and EOD
        difficulties_list = []
        for item in items:
            task = tasks.get(item["task_id"], {})
            note = task.get("note") or "N/A"
            difficulties_list.append(f"- {note}")
        difficulties_str = "\n".join(difficulties_list) if difficulties_list else "N/A"

        # ===== Báo cáo đầu ngày =====
        today_tasks = []
        for item in items:
            task = tasks.get(item["task_id"], {})
            percent = task.get("percent")
            percent_str = f"{percent}%" if percent is not None else "0-100%"
            today_tasks.append(f"- {item['task_name']}. `{percent_str}%`")
        today_tasks_str = "\n".join(today_tasks)

        morning_report = f"""**Daily Report < {date_display} > **
**Fresher: Trần Văn An**

--------------------------------------------------------------------------------------

Nội dung công việc ngày hôm nay:
{today_tasks_str}

--------------------------------------------------------------------------------------

Khó khăn:
{difficulties_str}

--------------------------------------------------------------------------------------

>Trello: https://trello.com/invite/b/6964c51e61b18697c47479cc/ATTIc1d3d8fe352d07f88ecb34cd8891b592A33ED4BB/fresher-dev
GitHub: https://github.com/antv-runs
Drive: https://drive.google.com/drive/folders/1d9RUUju0d78LPTJanfvLscxcBeZckvA_?usp=sharing
"""

        (daily_dir / "Báo cáo đầu ngày.txt").write_text(
            morning_report,
            encoding="utf-8",
        )

        # ===== Báo cáo cuối ngày =====
        done_tasks = []
        for item in items:
            task = tasks.get(item["task_id"], {})
            percent = task.get("percent")
            percent_str = f"{percent}%" if percent is not None else "100%"
            done_tasks.append(f"- {item['task_name']}. `{percent_str}%`")
        done_tasks_str = "\n".join(done_tasks)

        eod_report = f"""**Daily Report < {date_display} > **
**Fresher: Trần Văn An**

--------------------------------------------------------------------------------------

Nội dung công việc ngày hôm nay:
{done_tasks_str}

--------------------------------------------------------------------------------------

Nội dung công việc ngày mai:
{next_tasks}

--------------------------------------------------------------------------------------

Khó khăn:
{difficulties_str}

--------------------------------------------------------------------------------------

>Trello: https://trello.com/invite/b/6964c51e61b18697c47479cc/ATTIc1d3d8fe352d07f88ecb34cd8891b592A33ED4BB/fresher-dev
GitHub: https://github.com/antv-runs
Drive: https://drive.google.com/drive/folders/1d9RUUju0d78LPTJanfvLscxcBeZckvA_?usp=sharing
"""

        (daily_dir / "Báo cáo cuối ngày.txt").write_text(
            eod_report,
            encoding="utf-8",
        )

        # ===== Báo cáo DU2 =====
        du2_lines = [
            f"❖ Daily Report [{date_display}]",
            "TODAY:",
            "-----",
        ]

        for item in items:
            task = tasks.get(item["task_id"], {})

            estimate_h = task.get("estimate_h", "")
            actual_h = task.get("actual_h", "")
            note = task.get("note") or "N/A"
            percent = task.get("percent")
            percent_str = f"{percent}%" if percent is not None else "100%"

            du2_lines.append(f" -   {item['task_name']}. `{percent_str}%`")
            du2_lines.append(f"    - Time estimate: {estimate_h}h")
            du2_lines.append(f"    - Time complete: {actual_h}h")
            du2_lines.append(f"    - Difficulties at work: {note}")
            du2_lines.append("    - Lessons learned: N/A")  # TODO: Map from Excel if lesson column is added later
            du2_lines.append("")

        du2_lines.extend([
            "",
            "",
            "NEXT DAY:",
            "-----",
            next_tasks,
        ])

        (daily_dir / "Báo cáo DU2.txt").write_text(
            "\n".join(du2_lines),
            encoding="utf-8",
        )


# ========================
# Generate Word Report
# ========================
def generate_task_docx(tasks, daily_data):
    doc = Document()
    doc.add_heading("Quá trình thực hiện task", 0)

    grouped_by_task = {}

    for item in daily_data:
        grouped_by_task.setdefault(item["task_id"], []).append(item)

    for task_id, items in grouped_by_task.items():
        task = tasks.get(task_id, {})

        items = sorted(items, key=lambda x: x["date_obj"])

        task_name = task.get("task_name") or items[0]["task_name"]
        estimate_h = task.get("estimate_h", "")
        actual_h = task.get("actual_h", "")
        note = task.get("note") or "N/A"

        start_date = items[0]["date"]
        end_date = items[-1]["date"]

        doc.add_paragraph(f"Tên task: {task_name}")
        doc.add_paragraph(f"{start_date} ~ {end_date}")
        doc.add_paragraph(f"1. Nội dung task: {task_name}.")
        doc.add_paragraph(f"2. Estimated: {estimate_h}h")
        doc.add_paragraph(f"3. Time spent: {actual_h}h")
        doc.add_paragraph(f"4. Khó khăn/Lí do: {note}")
        doc.add_paragraph("5. Kinh nghiệm/Bài học rút ra: N/A")  # TODO: Map from Excel if lesson column is added later
        doc.add_paragraph("")

    doc.save(OUTPUT_DIR / "quá trình thực hiện task.docx")


# ========================
# Main
# ========================
def main():
    wb = load_workbook(INPUT_FILE, data_only=True)

    tasks = read_wbs(wb)
    daily_data = read_daily(wb)

    if not daily_data:
        print("Không có nội dung ngày nào trong sheet Report Daily.")
        return

    latest_date = max(item["date_obj"] for item in daily_data)
    latest_daily_data = [item for item in daily_data if item["date_obj"] == latest_date]

    generate_reports_by_date(latest_daily_data, tasks)
    generate_task_docx(tasks, daily_data)

    print("Done. Check folder Generate/")


if __name__ == "__main__":
    main()